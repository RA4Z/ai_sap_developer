import shutil
from docx import Document
from language_translation import Language
import re
import openpyxl
import xlwings as xw

def retornar_codigo(texto: str):
    indice_inicio = 0
    if indice_inicio == -1:
        return None

    indice_fim = texto.find("----fimpython----", indice_inicio)
    if indice_fim == -1:
        return None

    return texto[indice_inicio:indice_fim]


def pegar_texto_externo(texto: str):
    indice_fim = texto.find("----fimpython----")
    if indice_fim == -1:
        return None

    return texto[texto.find("\n", indice_fim) + 1:]

def copiar_pasta(origin: str, destiny: str):
    try:
        shutil.copytree(origin, destiny)
        print(f"Pasta copiada com sucesso de '{origin}' para '{destiny}'")
    except FileExistsError:
        print(f"A pasta '{destiny}' já existe. A cópia foi cancelada.")
    except Exception as e:
        print(f"Erro ao copiar pasta: {e}")

def cadastrar_doc(documentation: str, path: str):
    lang = Language()
    doc = Document()
    doc.add_paragraph(documentation)
    doc.save(f"{path}\\{lang.search('doc')}.docx")

def cadastrar_worksheet(output: str, path: str):
    match = re.search(r"ExcelHandler\('(.*?)'\)", output)
    if match:
        nome_arquivo = match.group(1)
        # Verifica a extensão do arquivo e cria a planilha com a biblioteca correta
        if nome_arquivo.endswith('.xlsm'):
            workbook = xw.Book()  # Cria com xlwings para macros
            workbook.sheets.add("Principal")  # Cria a sheet "Principal"
        else:
            workbook = openpyxl.Workbook()  # Cria com openpyxl para arquivos sem macros
            worksheet = workbook.active  # Pega a sheet ativa
            worksheet.title = "Principal"  # Define o nome da primeira planilha

        workbook.save(f'{path}/{nome_arquivo}')
        if nome_arquivo.endswith('.xlsm'):
            workbook.close()

        print(f"Planilha '{nome_arquivo}' criada com sucesso!")

